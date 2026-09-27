import re
import pymupdf
from docx import Document
from pptx import Presentation

from pathlib import Path

from app.schemas.documents import (
    DocumentExtractionResult,
    DocumentSegment,
)


def extract_txt(file_path: Path) -> DocumentExtractionResult:
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if file_path.suffix.lower() != ".txt":
        raise ValueError("Expected a TXT file")

    try:
        content = file_path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ValueError(
            "Unable to read TXT file using UTF-8 encoding"
        ) from exc

    segments: list[DocumentSegment] = []

    for line_number, line in enumerate(content.splitlines(), start=1):
        cleaned_line = line.strip()

        if not cleaned_line:
            continue

        segments.append(
            DocumentSegment(
                reference_type="line",
                reference_number=line_number,
                text=cleaned_line,
            )
        )

    return DocumentExtractionResult(
        filename=file_path.name,
        file_type="txt",
        total_segments=len(segments),
        full_text=content.strip(),
        segments=segments,
    )

MIN_NATIVE_TEXT_CHARACTERS = 80
MIN_NATIVE_TEXT_WORDS = 10
OCR_DPI = 300


def has_meaningful_pdf_text(text: str) -> bool:
    stripped = text.strip()

    if len(stripped) < MIN_NATIVE_TEXT_CHARACTERS:
        return False

    words = re.findall(
        r"[A-Za-z0-9]+",
        stripped,
    )

    return len(words) >= MIN_NATIVE_TEXT_WORDS


def extract_page_text_with_ocr(
    page,
    page_number: int,
) -> str:
    native_text = page.get_text(
        "text"
    ).strip()

    if has_meaningful_pdf_text(native_text):
        return native_text

    try:
        text_page = page.get_textpage_ocr(
            language="eng",
            dpi=OCR_DPI,
            full=True,
            tessdata=pymupdf.get_tessdata(),
        )

        return page.get_text(
            "text",
            textpage=text_page,
        ).strip()
    except Exception as exc:
        raise ValueError(
            f"OCR failed on PDF page {page_number}"
        ) from exc


def extract_pdf(file_path: Path) -> DocumentExtractionResult:
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if file_path.suffix.lower() != ".pdf":
        raise ValueError("Expected a PDF file")

    segments: list[DocumentSegment] = []
    full_text_parts: list[str] = []

    try:
        document = pymupdf.open(file_path)

        if document.needs_pass:
            raise ValueError("Password-protected PDF files are not supported")

        for page_number, page in enumerate(
            document,
            start=1,
        ):
            text = extract_page_text_with_ocr(
                page,
                page_number,
            )

            if not text:
                continue

            segments.append(
                DocumentSegment(
                    reference_type="page",
                    reference_number=page_number,
                    text=text,
                )
            )

            full_text_parts.append(text)

    except pymupdf.FileDataError as exc:
        raise ValueError("Invalid or corrupted PDF file") from exc

    finally:
        if "document" in locals():
            document.close()

    return DocumentExtractionResult(
        filename=file_path.name,
        file_type="pdf",
        total_segments=len(segments),
        full_text="\n\n".join(full_text_parts),
        segments=segments,
    )


def extract_docx(file_path: Path) -> DocumentExtractionResult:
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if file_path.suffix.lower() != ".docx":
        raise ValueError("Expected a DOCX file")

    try:
        document = Document(file_path)
    except Exception as exc:
        raise ValueError("Invalid or corrupted DOCX file") from exc

    segments: list[DocumentSegment] = []
    full_text_parts: list[str] = []

    for paragraph_number, paragraph in enumerate(
        document.paragraphs,
        start=1,
    ):
        text = paragraph.text.strip()

        if not text:
            continue

        segments.append(
            DocumentSegment(
                reference_type="paragraph",
                reference_number=paragraph_number,
                text=text,
            )
        )

        full_text_parts.append(text)

    return DocumentExtractionResult(
        filename=file_path.name,
        file_type="docx",
        total_segments=len(segments),
        full_text="\n\n".join(full_text_parts),
        segments=segments,
    )


def extract_pptx(file_path: Path) -> DocumentExtractionResult:
    if not file_path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")

    if file_path.suffix.lower() != ".pptx":
        raise ValueError("Expected a PPTX file")

    try:
        presentation = Presentation(file_path)
    except Exception as exc:
        raise ValueError("Invalid or corrupted PPTX file") from exc

    segments: list[DocumentSegment] = []
    full_text_parts: list[str] = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        slide_text_parts: list[str] = []

        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            if text:
                slide_text_parts.append(text)

        if not slide_text_parts:
            continue

        slide_text = "\n".join(slide_text_parts)

        segments.append(
            DocumentSegment(
                reference_type="slide",
                reference_number=slide_number,
                text=slide_text,
            )
        )

        full_text_parts.append(slide_text)

    return DocumentExtractionResult(
        filename=file_path.name,
        file_type="pptx",
        total_segments=len(segments),
        full_text="\n\n".join(full_text_parts),
        segments=segments,
    )

SUPPORTED_EXTENSIONS = {
    ".txt",
    ".pdf",
    ".docx",
    ".pptx",
}


def extract_document(file_path: Path) -> DocumentExtractionResult:
    extension = file_path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            f"Unsupported file type: {extension or 'unknown'}"
        )

    if extension == ".txt":
        return extract_txt(file_path)

    if extension == ".pdf":
        return extract_pdf(file_path)

    if extension == ".docx":
        return extract_docx(file_path)

    if extension == ".pptx":
        return extract_pptx(file_path)

    raise ValueError("Unsupported document type")